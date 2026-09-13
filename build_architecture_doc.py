from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = 'Bettersaid Technical Architecture.docx'
NAVY = '0E2841'
TEAL = '36BDB7'
LIGHT = 'EAF7F5'
BORDER = 'D9E4E8'

def set_font(run, size=None, bold=None, color='000000'):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for side in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + side
        elem = tcBorders.find(qn(tag))
        if elem is None:
            elem = OxmlElement(tag)
            tcBorders.append(elem)
        elem.set(qn('w:val'), 'single')
        elem.set(qn('w:sz'), '6')
        elem.set(qn('w:color'), BORDER)

def set_cell_text(cell, text, bold=False, color='000000'):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_font(run, 9.5, bold, color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    borders(cell)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)

title = doc.add_paragraph(style='Title')
title.paragraph_format.space_after = Pt(5)
run = title.add_run('Bettersaid Technical Architecture')
set_font(run, 27, True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run('Simple setup for transcript analysis and conversation practice')
set_font(run, 12, False, NAVY)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
run = p.add_run('Purpose. ')
set_font(run, 10.5, True)
run = p.add_run('Bettersaid helps people review sensitive conversations and practise for difficult discussions. The first version keeps the product simple: one website, one protected AI service, and a downloadable report.')
set_font(run, 10.5)

h = doc.add_paragraph(style='Heading 1')
h.paragraph_format.space_before = Pt(4)
h.paragraph_format.space_after = Pt(8)
set_font(h.add_run('Architecture at a glance'), 16, True)

table = doc.add_table(rows=1, cols=3)
table.autofit = False
table.columns[0].width = Inches(1.45)
table.columns[1].width = Inches(2.15)
table.columns[2].width = Inches(3.55)
headers = ['Part', 'Runs where', 'Responsibility']
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    shade(cell, NAVY)
    set_cell_text(cell, text, True, 'FFFFFF')
for row in [
    ('Bettersaid website', 'User browser', 'Lets users choose a mode, add a transcript, read the analysis, and download an HTML report.'),
    ('Protected AI function', 'Hosting platform', 'Receives the transcript, keeps the API key private, asks OpenAI for structured analysis, and returns the result.'),
    ('OpenAI Responses API', 'OpenAI platform', 'Creates neutral, structured insights based on Better Said instructions.'),
    ('Local HTML report', 'User computer', 'Lets the user save and share their final reflection without requiring a database.'),
]:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        shade(cells[i], LIGHT if len(table.rows) % 2 == 0 else 'FFFFFF')
        set_cell_text(cells[i], text, i == 0, NAVY if i == 0 else '000000')

h = doc.add_paragraph(style='Heading 1')
h.paragraph_format.space_before = Pt(18)
h.paragraph_format.space_after = Pt(7)
set_font(h.add_run('What happens step by step'), 16, True)

steps = [
    ('1', 'The user opens Bettersaid', 'They choose transcript reflection or conversation practice.'),
    ('2', 'The user adds their content', 'They paste or upload a transcript, or begin writing a practice message.'),
    ('3', 'The website sends a protected request', 'The browser sends the content to the app’s protected AI function. The API key never appears in the browser.'),
    ('4', 'The protected function calls OpenAI', 'It sends the Better Said instructions and the conversation content to the OpenAI Responses API.'),
    ('5', 'OpenAI creates structured analysis', 'The response contains an overview, direct quotations, communication patterns, possible effects, and practical ways forward.'),
    ('6', 'Bettersaid displays the result', 'The user reads the analysis in the report screen and can continue the reflection through conversation practice.'),
    ('7', 'The user saves the result', 'Bettersaid creates a local HTML report that the user can keep or share.'),
]
for num, heading, body in steps:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.05)
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(num + '  ')
    set_font(r, 11, True, TEAL)
    r = p.add_run(heading + '. ')
    set_font(r, 10.5, True, NAVY)
    r = p.add_run(body)
    set_font(r, 10.5)

h = doc.add_paragraph(style='Heading 1')
h.paragraph_format.space_before = Pt(14)
h.paragraph_format.space_after = Pt(7)
set_font(h.add_run('Privacy and security'), 16, True)
for label, body in [
    ('API key', 'Stored only as a protected environment variable on the hosting platform. It is never placed in browser code.'),
    ('Transcript storage', 'The simple product version processes transcripts for analysis and does not require a permanent user database.'),
    ('Saved output', 'The user chooses whether to download a local HTML report.'),
]:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + ': ')
    set_font(r, 10.5, True, NAVY)
    r = p.add_run(body)
    set_font(r, 10.5)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(0)
r = p.add_run('One sentence summary: ')
set_font(r, 10.5, True, NAVY)
r = p.add_run('Bettersaid is a secure web app that turns a conversation into a neutral AI reflection and lets the user save the result locally.')
set_font(r, 10.5)

doc.save(OUT)
